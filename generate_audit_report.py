from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "DSLite项目代码审计_Debug过程报告_2026-05-29更新.xlsx"
DATE = "2026-05-29"
COLORS = {
    "navy": "1F4E78",
    "green": "E2F0D9",
    "yellow": "FFF2CC",
    "red": "FCE4D6",
    "gray": "F2F2F2",
    "white": "FFFFFF",
}


def style(ws, widths):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    for cell in ws[1]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=COLORS["white"])
        cell.fill = PatternFill("solid", fgColor=COLORS["navy"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 26


def add_sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style(ws, widths)
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    add_sheet(wb, "报告概览", ["项目", "内容"], [
        ["项目路径", str(ROOT)],
        ["报告更新日期", DATE],
        ["报告生成时间", now],
        ["审计对象", "DSLite F28034 烧录调试上位机：v2_改进版/dslite_tool_v2.py、发布目录、随包 DSLite/DebugServer、README。"],
        ["总体结论", "项目已完成代码审计、发布打包、硬件非破坏性联调和用户反馈修复。当前主线为 v2_改进版，桌面发布目录为 DSLite_F28034_Tool_Release。"],
        ["当前发布状态", "exe 可启动；_internal 依赖目录已补齐；后台 DSLite 黑框已隐藏；程序内已禁止并发 DSLite 任务。"],
        ["硬件验证状态", "XDS100v2/F28034 已连接验证。flash --list-ops、flash --list-settings、memory 读取 DEVEMU 寄存器均通过。未执行擦除、解锁或烧录。"],
        ["本次功能改进", "增加固定代码目录、最近文件路径保存、CSM 密码保存/加载，以及“写入密码+上锁”功能。"],
    ], [24, 120])

    findings = [
        ["F-001", DATE, "严重", "旧版硬编码 DSLite 路径", "dslite_tool.py:50", "换目录或换机器后直接失效", "保留旧版归档，实际使用 v2_改进版", "已处理"],
        ["F-002", DATE, "严重", "相对 ccxml 从 DSLite 工作目录解析，容易找错文件", "dslite_tool_v2.py", "命令可能找不到 f28034.ccxml", "新增 resolve_ccxml_path，统一按脚本或 exe 所在目录解析", "已修复"],
        ["F-003", DATE, "严重", "发布版 exe 打不开，提示找不到指定模块", "发布目录", "PyInstaller onedir 漏复制 _internal 依赖目录", "补齐 _internal 并验证 GUI exe 能启动", "已修复"],
        ["F-004", DATE, "严重", "每次点击功能弹出黑色控制台窗口", "CmdThread.run", "用户体验差", "subprocess.run 增加 CREATE_NO_WINDOW", "已修复"],
        ["F-005", DATE, "严重", "XDS100/DSLite 并发会话导致 Error -151", "run_cmd", "并行启动两个 DSLite 命令时一个会失败", "程序内禁止并发任务，已有任务时提示等待", "已修复"],
        ["F-006", DATE, "中", "读取芯片ID命令使用不存在的 GetDeviceID 操作", "read_device_id", "按钮功能必失败", "改为读取 DEVEMU 寄存器 0x0880,0x08", "已修复"],
        ["F-007", DATE, "中", "检测设备命令不符合当前 DSLite 操作集合", "check_device_status", "Connect/GetDeviceID 不在 flash 操作列表中", "改为非破坏性的 flash --list-ops", "已修复"],
        ["F-008", DATE, "中", "parse_bin.py 异常分支引用不存在变量 file", "parse_bin.py:24", "文件不存在时二次报错", "改为 file_name", "已修复"],
        ["F-009", DATE, "中", "README 和历史报告脚本存在编码乱码", "README.md、v2_改进版/README.md、generate_audit_report.py", "维护和交付阅读困难", "重写 README，重新生成中文 Excel 报告", "已修复"],
        ["F-010", DATE, "中", "程序文件没有固定存放区", "Flash 烧录参数", "每次选择程序文件都要重新定位目录", "增加固定代码目录配置，文件选择默认从该目录打开", "已修复"],
        ["F-011", DATE, "中", "CSM 解锁/上锁密码无法持久保存", "CSM 密码设置", "解锁和上锁操作容易手工输错", "增加 CSM Key0~Key7 保存/加载，写入密码+上锁前二次确认", "已修复"],
        ["F-012", DATE, "低", "仓库包含大量 TI 工具链二进制或未跟踪文件", "test/ccs_base", "Git 状态噪声大，仓库体积膨胀", "建议后续区分源码仓库和发布包", "待处理"],
    ]
    ws = add_sheet(wb, "问题清单", ["编号", "日期", "级别", "问题", "位置", "影响", "处理或建议", "状态"], findings, [12, 14, 10, 34, 34, 46, 58, 12])
    for row in ws.iter_rows(min_row=2):
        row[2].fill = PatternFill("solid", fgColor=COLORS["red"] if row[2].value == "严重" else COLORS["yellow"] if row[2].value == "中" else COLORS["gray"])
        if row[7].value in ("已修复", "已处理"):
            row[7].fill = PatternFill("solid", fgColor=COLORS["green"])

    add_sheet(wb, "Debug过程", ["步骤", "日期", "阶段", "执行动作", "结果或证据"], [
        [1, DATE, "定位项目", "扫描桌面 2026-4-2 项目目录", "确认项目路径 C:\\Users\\hjs\\Desktop\\test--2026-4-2"],
        [2, DATE, "代码审计", "检查 subprocess、ccxml、DSLite、flash、memory、erase、异常处理", "定位硬编码路径、配置同步、命令格式、异常处理等问题"],
        [3, DATE, "基础修复", "修复 ccxml 路径、DSLite.exe 绝对路径、随包工具链检测、parse_bin 异常变量", "py_compile 和 unittest 通过"],
        [4, DATE, "首次打包", "使用 PyInstaller onedir 生成 DSLite_F28034_Tool.exe", "生成发布目录 DSLite_F28034_Tool_Release"],
        [5, DATE, "发布问题修复", "用户反馈 exe 找不到指定模块", "确认漏复制 _internal，补齐后 GUI exe 启动成功"],
        [6, DATE, "硬件联调", "连接 XDS100v2/F28034，执行 help、list-ops、list-settings、memory", "硬件链路通过，DEVEMU 读取成功"],
        [7, DATE, "并发问题定位", "并行执行两个 DSLite 命令", "复现 Error -151，确认 XDS100/DSLite 不适合并发会话"],
        [8, DATE, "功能修复", "检测设备改为 flash --list-ops，读取芯片ID改为 memory 0x0880,0x08", "避免调用不存在的 GetDeviceID 操作"],
        [9, DATE, "体验修复", "隐藏后台 DSLite 控制台窗口", "加入 CREATE_NO_WINDOW，点击功能不再弹黑框"],
        [10, DATE, "功能改进", "增加固定代码目录、最近文件、CSM 密码保存和写入密码+上锁", "配置保存到 dslite_config.json，上锁操作增加二次确认"],
        [11, DATE, "最终发布", "重新打包并覆盖发布目录 exe 和 _internal", "GUI exe 可启动，发布目录完整"],
    ], [10, 14, 18, 54, 76])

    add_sheet(wb, "修复清单", ["日期", "文件或目录", "修改", "目的"], [
        [DATE, "v2_改进版/dslite_tool_v2.py", "新增冻结态资源目录识别，打包后按 exe 所在目录查找资源", "发布目录可独立运行"],
        [DATE, "v2_改进版/dslite_tool_v2.py", "run_cmd 增加已有任务检查", "防止并发 DSLite 会话导致 XDS100/FTDI 错误"],
        [DATE, "v2_改进版/dslite_tool_v2.py", "subprocess.run 增加 CREATE_NO_WINDOW", "隐藏每次点击功能弹出的黑色控制台窗口"],
        [DATE, "v2_改进版/dslite_tool_v2.py", "check_device_status 改为 flash --list-ops", "使用当前 DSLite 支持的非破坏性连接检测"],
        [DATE, "v2_改进版/dslite_tool_v2.py", "read_device_id 改为 memory --range=0x0880,0x08", "读取 DEVEMU/芯片标识寄存器"],
        [DATE, "桌面发布目录", "补齐 _internal 依赖目录", "修复 exe 找不到指定模块"],
        [DATE, "README.md、v2_改进版/README.md", "重写中文文档并标注日期", "记录当前发布和硬件验证状态"],
        [DATE, "v2_改进版/dslite_tool_v2.py", "新增 code_dir、last_load_file、last_elf_file、last_bin_file 配置", "建立固定代码存放区并记住最近程序文件"],
        [DATE, "v2_改进版/dslite_tool_v2.py", "新增 csm_keys 配置和保存按钮", "保存解锁/上锁密码，减少手工输入错误"],
        [DATE, "v2_改进版/dslite_tool_v2.py", "新增写入密码+上锁按钮", "执行 ProgramPassword 和 Lock，并在执行前二次确认"],
    ], [14, 36, 62, 64])

    add_sheet(wb, "验证记录", ["日期", "验证命令或动作", "结果", "说明"], [
        [DATE, "python -m py_compile v2_改进版\\dslite_tool_v2.py", "通过", "源码语法检查通过"],
        [DATE, "python -m unittest tests.test_dslite_tool_v2_helpers", "通过", "4 项单元测试 OK"],
        [DATE, "启动桌面发布版 DSLite_F28034_Tool.exe", "通过", "进程保持运行，GUI 可启动"],
        [DATE, "DSLite.exe help", "通过", "输出 DSLite version 7.2.0.1988"],
        [DATE, "DSLite.exe flash --config=f28034.ccxml --list-ops --verbose", "通过", "返回 Erase、CalculateChecksum、DepletionRecovery、FrequencyTest、ProgramPassword、Lock、Unlock"],
        [DATE, "DSLite.exe flash --config=f28034.ccxml --list-settings --verbose", "通过", "串行执行成功"],
        [DATE, "DSLite.exe memory --config=f28034.ccxml --range=0x0880,0x08 --output=device_id_verify.bin --verbose", "通过", "读取结果：6B 0F BB 01 00 00 F7 FF"],
        [DATE, "并行执行两个 DSLite 命令", "复现问题", "其中一个会报 Error -151，程序已改为禁止并发任务"],
        [DATE, "python -m py_compile v2_改进版\\dslite_tool_v2.py", "通过", "固定代码目录和 CSM 密码功能加入后语法检查通过"],
    ], [14, 78, 16, 76])

    add_sheet(wb, "发布说明", ["日期", "项目", "内容"], [
        [DATE, "发布目录", "C:\\Users\\hjs\\Desktop\\DSLite_F28034_Tool_Release"],
        [DATE, "必须保留", "DSLite_F28034_Tool.exe、_internal\\、f28034.ccxml、test\\ccs_base\\、dslite_config.json"],
        [DATE, "打包方式", "PyInstaller onedir；不能只复制 exe，必须保留 _internal"],
        [DATE, "硬件工具链", "随包 DSLite 7.2.0.1988 / DebugServer"],
        [DATE, "配置文件", "dslite_config.json 保存固定代码目录、最近文件路径和 CSM Key0~Key7"],
        [DATE, "未执行操作", "未执行擦除、解锁或烧录，未修改芯片 Flash 内容"],
    ], [14, 24, 110])

    add_sheet(wb, "残余风险", ["日期", "风险", "说明", "建议"], [
        [DATE, "擦除和解锁风险", "Unlock、Erase 属于破坏性操作，误点可能清空芯片", "增加二次确认，显示芯片、ccxml、目标地址摘要"],
        [DATE, "BIN 地址校验不足", "当前只检查非空，没有校验 F28034 Flash 范围和扇区边界", "增加地址范围、扇区边界、文件大小越界校验"],
        [DATE, "发布包体积大", "随包 test/ccs_base 约 400MB 以上", "后续可整理最小运行集或改为检测系统 CCS 安装"],
        [DATE, "源码仓库噪声", "大量 TI 二进制文件进入 Git 状态", "区分源码仓库与发布包，完善 .gitignore"],
        [DATE, "CSM 密码保管", "密码保存在本机 dslite_config.json", "将配置文件纳入本地备份，不要随意共享包含真实密码的配置"],
    ], [14, 30, 62, 62])

    add_sheet(wb, "后续建议", ["日期", "优先级", "事项", "说明"], [
        [DATE, 1, "补安全确认", "对擦除、解锁、烧录 BIN 增加二次确认和参数摘要"],
        [DATE, 2, "补地址校验", "限定 F28034 Flash 合法区间、扇区起始地址，BIN 文件大小不能跨越非法区域"],
        [DATE, 3, "补密码备份策略", "增加导入/导出 CSM 密码配置，避免密码丢失"],
        [DATE, 4, "整理发布包", "研究最小 DSLite 运行依赖，减少发布目录体积"],
        [DATE, 5, "整理版本主线", "明确以 v2_改进版 为主线，旧版只保留归档"],
        [DATE, 6, "完善测试文档", "记录接线、供电、XDS100 驱动、常见 Error -151 处理方式"],
    ], [14, 10, 30, 90])

    wb.save(OUT)
    loaded = load_workbook(OUT)
    assert loaded["报告概览"]["A1"].value == "项目"
    assert loaded["验证记录"]["D8"].value == "读取结果：6B 0F BB 01 00 00 F7 FF"
    print(OUT)


if __name__ == "__main__":
    main()
